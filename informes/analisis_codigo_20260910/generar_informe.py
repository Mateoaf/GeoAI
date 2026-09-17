"""Compone el informe y exige una explicación para cada celda real de código."""
from pathlib import Path
import ast
import hashlib
import html
import json
import re
import unicodedata

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[1]
STEM = 'INFORME_ANALISIS_CODIGO_Y_NOTEBOOKS'


def slug(text):
    value = ''.join(c for c in unicodedata.normalize('NFKD', text.lower()) if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]+', '-', value).strip('-')


def inline(value):
    # Escape all input; permit only links and elementary Markdown formatting.
    parts = re.split(r'(`[^`]+`|\*\*.+?\*\*|\[[^\]]+\]\([^)]+\))', value)
    out = []
    for part in parts:
        if part.startswith('`') and part.endswith('`'):
            out.append('<code>' + html.escape(part[1:-1]) + '</code>')
        elif part.startswith('**') and part.endswith('**'):
            out.append('<strong>' + inline(part[2:-2]) + '</strong>')
        else:
            match = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', part)
            if match:
                url = match[2]
                if url.startswith(('javascript:', 'data:')):
                    raise ValueError('Unexpected URL')
                out.append('<a href="'+html.escape(url, quote=True)+'">'+html.escape(match[1])+'</a>')
            else:
                out.append(html.escape(part))
    return ''.join(out)


def blocks(markdown):
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith('```'):
            language = line[3:]
            content = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                content.append(lines[i]); i += 1
            if i == len(lines):
                raise ValueError('Unclosed code fence')
            yield 'code', language, '\n'.join(content)
        elif line.startswith('#'):
            match = re.match(r'^(#{1,6}) (.+)$', line)
            if not match:
                raise ValueError(line)
            yield 'heading', len(match[1]), match[2]
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                row = [v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', v) for v in row):
                    rows.append(row)
                i += 1
            yield 'table', rows, None
            continue
        elif re.match(r'^\d+\. ', line):
            items = []
            while i < len(lines) and re.match(r'^\d+\. ', lines[i]):
                items.append(re.sub(r'^\d+\. ', '', lines[i])); i += 1
            yield 'list', items, None
            continue
        else:
            content = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and not lines[i].startswith(('#', '|', '```')) and not re.match(r'^\d+\. ', lines[i]):
                content.append(lines[i]); i += 1
            yield 'paragraph', ' '.join(content), None
            continue
        i += 1


def render_html(markdown):
    rendered, navigation = [], []
    for kind, a, b in blocks(markdown):
        if kind == 'heading':
            key = slug(b)
            rendered.append(f'<h{a} id="{key}">{inline(b)}</h{a}>')
            if a == 2:
                navigation.append(f'<a href="#{key}">{html.escape(b)}</a>')
        elif kind == 'paragraph':
            rendered.append('<p>'+inline(a)+'</p>')
        elif kind == 'list':
            rendered.append('<ol>'+''.join('<li>'+inline(x)+'</li>' for x in a)+'</ol>')
        elif kind == 'code':
            code='<pre><code>'+html.escape(b)+'</code></pre>'
            if a == 'python':
                code='<details class="source"><summary>Ver código original de la celda</summary>'+code+'</details>'
            rendered.append(code)
        elif kind == 'table':
            rendered.append('<div class="table"><table><thead><tr>'+''.join('<th>'+inline(x)+'</th>' for x in a[0])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+inline(x)+'</td>' for x in row)+'</tr>' for row in a[1:])+'</tbody></table></div>')
    style = '''
:root{--ink:#183342;--accent:#246b64;--paper:#fff;--muted:#65727b}*{box-sizing:border-box}body{margin:0;color:var(--ink);background:#f3f5f4;font:16px/1.7 system-ui,Segoe UI,sans-serif}aside{position:fixed;inset:0 auto 0 0;width:290px;overflow:auto;background:#153c3b;color:#fff;padding:24px 20px}aside h2{margin:0 0 18px;font-size:20px;color:#fff}aside a{display:block;font-size:13px;line-height:1.4;color:#e4f2ec;text-decoration:none;padding:7px 0;border-bottom:1px solid #ffffff20}aside a:hover{color:#fff;text-decoration:underline}main{margin-left:290px;max-width:1280px;padding:35px 55px 100px;background:var(--paper)}h1{font-size:36px;line-height:1.2;margin-top:0}h2{font-size:26px;line-height:1.35;margin:55px 0 20px;border-top:2px solid #d6e4df;padding-top:22px;scroll-margin-top:20px}h3{font-size:20px;margin-top:32px}h4{font-size:17px;margin:28px 0 12px;padding:12px 15px;background:#edf4f1;border-left:4px solid var(--accent)}a{color:var(--accent)}p{margin:14px 0}code{font:0.87em Consolas,monospace;background:#eef2f1;padding:2px 4px;border-radius:3px;overflow-wrap:anywhere}pre{background:#132d38;color:#edf5f6;padding:18px;overflow:auto;font:12.5px/1.55 Consolas,monospace;border-radius:5px}pre code{background:none;padding:0;color:inherit;overflow-wrap:normal;white-space:pre}summary{cursor:pointer;font-size:14px;font-weight:600;color:var(--accent);padding:9px 12px;background:#f0f5f3;border:1px solid #d4e3dd;border-radius:5px}.source{margin:12px 0}table{border-collapse:collapse;font-size:13px;line-height:1.5;width:100%}th,td{border:1px solid #d8e3df;padding:9px;text-align:left;vertical-align:top}th{background:#e4eeea}tr:nth-child(even){background:#f7f9f8}.table{overflow:auto;margin:24px 0}.toolbar{display:flex;gap:8px;flex-wrap:wrap;margin:0 0 24px}button{padding:9px 14px;border:1px solid #9ebbb1;background:#f2f7f4;color:#194f46;border-radius:5px;cursor:pointer}.hint{color:var(--muted);font-size:13px}li{margin:10px 0}@media(max-width:950px){aside{position:static;width:auto;max-height:290px}main{margin:0;padding:28px 20px}h1{font-size:29px}}@media print{aside,.toolbar,.hint{display:none}main{margin:0;padding:0;max-width:none}body{font-size:10pt;background:#fff}h2{break-before:page}pre{white-space:pre-wrap;overflow-wrap:anywhere;font-size:7pt;background:#f4f4f4;color:#111}pre code{white-space:pre-wrap}h3,h4{break-after:avoid}tr{break-inside:avoid}a{color:inherit;text-decoration:none}.table{overflow:visible}}
'''
    return '<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>GeoAI — Informe de código y notebooks</title><style>'+style+'</style></head><body><aside><h2>GeoAI · revisión técnica</h2>'+''.join(navigation)+'</aside><main><div class="toolbar"><button onclick="document.querySelectorAll(\'details\').forEach(x=>x.open=true)">Mostrar todo el código</button><button onclick="document.querySelectorAll(\'details\').forEach(x=>x.open=false)">Ocultar código</button><button onclick="window.print()">Imprimir / PDF</button></div><p class="hint">15 notebooks · 95 celdas · 44 pasos del plan. Usa Ctrl+F para buscar. El informe funciona sin conexión.</p>'+''.join(rendered)+'</main><script>window.addEventListener("beforeprint",()=>document.querySelectorAll("details").forEach(x=>x.open=true));</script></body></html>'


def render_docx(markdown):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    doc = Document()
    sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Cm(1.9)
    sec.left_margin=sec.right_margin=Cm(2)
    normal=doc.styles['Normal']
    normal.font.name='Calibri'; normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(7)
    for n in range(1,5):
        doc.styles[f'Heading {n}'].font.color.rgb=RGBColor.from_string('246B64')
    header=sec.header.paragraphs[0]
    header.text='GEOAI  |  Análisis de código y prospectividad aurífera  |  10/09/2026'
    header.runs[0].font.size=Pt(8)
    footer=sec.footer.paragraphs[0]
    footer.alignment=2
    footer.add_run('Página ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    def add_inline(paragraph, text):
        for part in re.split(r'(`[^`]+`|\*\*.+?\*\*|\[[^\]]+\]\([^)]+\))',text):
            if part.startswith('`') and part.endswith('`'):
                r=paragraph.add_run(part[1:-1]);r.font.name='Consolas';r.font.size=Pt(8)
            elif part.startswith('**') and part.endswith('**'):
                paragraph.add_run(part[2:-2]).bold=True
            else:
                m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part)
                if m:
                    relation=paragraph.part.relate_to(m[2],RELATIONSHIP_TYPE.HYPERLINK,is_external=True)
                    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),relation)
                    run=OxmlElement('w:r');properties=OxmlElement('w:rPr')
                    color=OxmlElement('w:color');color.set(qn('w:val'),'246B64');properties.append(color)
                    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');properties.append(underline)
                    run.append(properties);text=OxmlElement('w:t');text.text=m[1];run.append(text)
                    link.append(run);paragraph._p.append(link)
                else:
                    paragraph.add_run(part)
    for kind,a,b in blocks(markdown):
        if kind=='heading':
            if a==1: doc.add_heading(b,0)
            else:
                p=doc.add_heading(b,min(a-1,4))
                if a==2:p.paragraph_format.page_break_before=True
        elif kind=='paragraph':add_inline(doc.add_paragraph(),a)
        elif kind=='list':
            for item in a:add_inline(doc.add_paragraph(style='List Number'),item)
        elif kind=='code':
            for line in (b.splitlines() or ['# Celda vacía']):
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1
                r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(7)
        elif kind=='table':
            table=doc.add_table(rows=0, cols=len(a[0]));table.style='Light Shading Accent 1'
            for row in a:
                cells=table.add_row().cells
                for cell,value in zip(cells,row):add_inline(cell.paragraphs[0],value)
    doc.core_properties.title='Análisis técnico del código y notebooks de prospectividad aurífera'
    doc.core_properties.subject='GeoAI: 15 notebooks, 95 celdas y contraste con el plan de 44 pasos'
    doc.core_properties.author='Revisión técnica del proyecto GeoAI'
    doc.save(BASE/(STEM+'.docx'))


def main():
    notebooks={}
    for p in sorted((ROOT/'notebooks').glob('[0-9][0-9]_*.ipynb')):
        notebooks[p.name[:2]]=(p,json.loads(p.read_text(encoding='utf-8')))
    expected={(n,i) for n,(_,book) in notebooks.items() for i,c in enumerate(book['cells']) if c['cell_type']=='code'}
    visited=[]
    sources=[BASE/name for name in ('01_contexto.md','02_notebooks_a_c.md','03_notebooks_d.md','04_notebooks_e_f.md','05_resultados_y_revision.md')]
    content='\n\n'.join(p.read_text(encoding='utf-8') for p in sources)
    def replace_cell(match):
        num,index,title=match[1],int(match[2]),match[3]
        path,book=notebooks[num]
        cell=book['cells'][index]
        assert cell['cell_type']=='code'
        visited.append((num,index))
        source=''.join(cell['source'])
        # Source text is copied, never executed.
        ast.parse(source)
        marker=cell.get('execution_count')
        header=f'#### NB{num} · Celda {index} — {title}'
        info=f'Posición {index+1} del cuaderno; contador guardado: {marker if marker is not None else "sin contador"}. Fuente: [{path.name}](../../notebooks/{path.name}).'
        code=source.rstrip() if source.strip() else '# Celda original vacía; no contiene instrucciones.'
        return header+'\n\n'+info+'\n\n```python\n'+code+'\n```'
    content=re.sub(r'^@@CELL (\d{2}) (\d+) (.+)$',replace_cell,content,flags=re.M)
    if len(visited)!=len(set(visited)) or set(visited)!=expected:
        raise ValueError({'missing':sorted(expected-set(visited)),'extra':sorted(set(visited)-expected),'duplicates':len(visited)-len(set(visited))})
    inventory=['| Notebook | Fase | Celdas totales | Celdas de código | Celdas con instrucciones ejecutables |','|---|---|---:|---:|---:|']
    for n,(p,book) in notebooks.items():
        code=[c for c in book['cells'] if c['cell_type']=='code']
        phase='A' if n=='00' else 'B' if n=='01' else 'C' if n=='02' else 'D' if int(n)<=6 else 'E' if int(n)<=9 else 'F'
        active=sum(bool(ast.parse(''.join(c['source'])).body) for c in code)
        inventory.append(f'| `{p.stem}` | {phase} | {len(book["cells"])} | {len(code)} | {active} |')
    content=content.replace('{{INVENTARIO}}','\n'.join(inventory))
    if '@@CELL' in content or '{{' in content:raise ValueError('Unresolved template')
    (BASE/(STEM+'.md')).write_text(content,encoding='utf-8')
    (BASE/(STEM+'.html')).write_text(render_html(content),encoding='utf-8')
    render_docx(content)
    paragraphs=[]
    for m in re.finditer(r'^#### NB(\d{2}) · Celda (\d+) — (.+)$',content,re.M):
        paragraphs.append({'notebook':m[1],'cell_index_0':int(m[2]),'title':m[3],'anchor':slug(m[0][5:])})
    evidence=dict(notebooks=len(notebooks),cells_expected=len(expected),cells_explained=len(visited),
        all_cells_covered=True,approximate_word_count=len(content.split()),
        cells=paragraphs,files=[{'name':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(BASE.glob(STEM+'.*'))])
    (BASE/'control_informe.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in evidence.items() if k not in ('cells','files')},ensure_ascii=False))
    print('\n'.join(x['name']+' '+str(x['bytes'])+' bytes' for x in evidence['files']))


if __name__=='__main__':main()
