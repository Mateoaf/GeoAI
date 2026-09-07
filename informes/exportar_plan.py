"""Exporta el plan Markdown a DOCX y una lista de tareas CSV; no altera fuentes GIS."""
from pathlib import Path
import csv
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
source = ROOT / 'PLAN_PROSPECTIVIDAD_AURIFERA_ESPANA.md'
text = source.read_text(encoding='utf-8')

tasks = []
phase = ''
for line in text.splitlines():
    if line.startswith('### Fase '):
        phase = line.removeprefix('### ')
    match = re.match(r'\*\*(\d{2})\. (.*?)\*\*(.*)', line)
    if match:
        num, title, detail = match.groups()
        tasks.append({'paso': num, 'fase': phase, 'tarea': title.rstrip('.'),
                      'descripcion_inicial': detail.strip(), 'estado': 'Pendiente',
                      'responsable': '', 'fecha_inicio': '', 'fecha_fin': '',
                      'evidencia_entrega': ''})
assert [int(t['paso']) for t in tasks] == list(range(1, 45))
with (ROOT / 'PLAN_TAREAS_PROSPECTIVIDAD.csv').open('w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=list(tasks[0]), delimiter=';')
    writer.writeheader()
    writer.writerows(tasks)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.7)
section.bottom_margin = Inches(.7)
section.left_margin = Inches(.7)
section.right_margin = Inches(.7)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(6)
for name in ['Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[name].font.color.rgb = RGBColor.from_string('203C52')

def hyperlink(paragraph, label, url):
    rel_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    el = OxmlElement('w:hyperlink')
    el.set(qn('r:id'), rel_id)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '1565A0')
    props.append(color)
    run.append(props)
    t = OxmlElement('w:t')
    t.text = label
    run.append(t)
    el.append(run)
    paragraph._p.append(el)

def inline(paragraph, value):
    pattern = r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)'
    for piece in re.split(pattern, value):
        if not piece:
            continue
        link = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', piece)
        if link:
            hyperlink(paragraph, link[1], link[2])
        elif piece.startswith('**') and piece.endswith('**'):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith('`') and piece.endswith('`'):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        else:
            paragraph.add_run(piece)

lines = text.splitlines()
i = 0
in_code = False
while i < len(lines):
    line = lines[i]
    if line.startswith('```'):
        in_code = not in_code
        i += 1
        continue
    if in_code:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(8)
    elif line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            row = [v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?', v) for v in row):
                rows.append(row)
            i += 1
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.style = 'Light Shading Accent 1'
        for ri, row in enumerate(rows):
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                inline(cell.paragraphs[0], value)
            if ri == 0:
                repeat = OxmlElement('w:tblHeader')
                table.rows[-1]._tr.get_or_add_trPr().append(repeat)
        doc.add_paragraph()
        continue
    elif line.startswith('# '):
        doc.add_heading(line[2:], level=0)
    elif line.startswith('## '):
        doc.add_heading(line[3:], level=1)
    elif line.startswith('### '):
        doc.add_heading(line[4:], level=2)
    elif line.strip():
        p = doc.add_paragraph()
        if re.match(r'\*\*\d{2}\.', line):
            p.paragraph_format.keep_with_next = True
        inline(p, line)
    i += 1

footer = section.footer.paragraphs[0]
footer.add_run('Plan de prospectividad aurífera · 05/09/2026 · ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
dest = ROOT / 'PLAN_PROSPECTIVIDAD_AURIFERA_ESPANA.docx'
doc.save(dest)
check = Document(dest)
expected_tables = sum(line.startswith('|') and (i == 0 or not lines[i-1].startswith('|')) for i, line in enumerate(lines))
assert len(check.tables) == expected_tables, len(check.tables)
for task in tasks:
    prefix = task['paso'] + '. ' + task['tarea']
    assert sum(p.text.startswith(prefix) for p in check.paragraphs) == 1, prefix
print(f'Plan exportado: {len(tasks)} tareas, {len(check.tables)} tablas, {len(check.paragraphs)} párrafos.')
