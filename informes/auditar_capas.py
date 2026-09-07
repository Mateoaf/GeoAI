"""Auditoría de solo lectura; escribe resultados únicamente en informes/auditoria_2026-09-05."""
from pathlib import Path
import json, sqlite3, zipfile, xml.etree.ElementTree as ET
import numpy as np
import rasterio
from pypdf import PdfReader
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'informes' / 'auditoria_2026-09-05'
OUT.mkdir(exist_ok=True)
def save(name, obj):
    (OUT/name).write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
for p in list((ROOT/'informes').glob('*.pdf')) + list(ROOT.glob('*.pdf')):
    reader = PdfReader(p)
    (OUT/(p.stem+'.txt')).write_text('\n\n'.join(f'--- PÁGINA {i+1} ---\n'+(page.extract_text() or '') for i,page in enumerate(reader.pages)), encoding='utf-8')
for p in (ROOT/'informes').glob('*.docx'):
    doc=Document(p)
    parts=[x.text for x in doc.paragraphs]
    parts += ['\n'.join(' | '.join(c.text for c in r.cells) for r in t.rows) for t in doc.tables]
    (OUT/(p.stem+'.txt')).write_text('\n'.join(parts),encoding='utf-8')
for p in (ROOT/'informes').glob('*.ipynb'):
    nb=json.loads(p.read_text(encoding='utf-8'))
    (OUT/(p.stem+'.txt')).write_text('\n\n'.join(f'--- CELDA {i} {c["cell_type"]} ---\n'+''.join(c.get('source',[]))+'\nOUTPUT:\n'+str(c.get('outputs',[])) for i,c in enumerate(nb['cells'])),encoding='utf-8')
inventory=[]
for p in ROOT.rglob('*'):
    if p.is_file() and OUT not in p.parents:
        inventory.append({'path':str(p.relative_to(ROOT)),'bytes':p.stat().st_size})
save('archivos.json',inventory)
gp=[]
for p in ROOT.glob('*.gpkg'):
    con=sqlite3.connect(p.as_uri()+'?mode=ro',uri=True); con.row_factory=sqlite3.Row
    item={'file':p.name,'tables':[]}
    for row in con.execute('SELECT * FROM gpkg_contents'):
        d=dict(row); name=d['table_name']; quoted='"'+name.replace('"','""')+'"'
        d['count']=con.execute(f'SELECT COUNT(*) FROM {quoted}').fetchone()[0]
        d['columns']=[dict(r) for r in con.execute(f'PRAGMA table_info({quoted})')]
        d['samples']=[{k:v for k,v in dict(r).items() if not isinstance(v,bytes)} for r in con.execute(f'SELECT * FROM {quoted} LIMIT 3')]
        item['tables'].append(d)
    item['geometry_columns']=[dict(r) for r in con.execute('SELECT * FROM gpkg_geometry_columns')]
    gp.append(item);con.close()
save('geopackages.json',gp)
rr=[]
for p in ROOT.glob('*.tif'):
    with rasterio.open(p) as src:
        d={'file':p.name,'crs':str(src.crs),'shape':[src.height,src.width],'bounds':list(src.bounds),'res':src.res,'nodata':src.nodata,'descriptions':src.descriptions,'tags':src.tags(),'bands':[]}
        for b in range(1,src.count+1):
            a=src.read(b,masked=True); vals=a.compressed(); vals=vals[np.isfinite(vals)]; unique=np.unique(vals)
            d['bands'].append({'band':b,'dtype':src.dtypes[b-1],'valid':len(vals),'min':float(vals.min()) if len(vals) else None,'max':float(vals.max()) if len(vals) else None,'unique_count':len(unique),'values':unique.tolist() if len(unique)<30 else None,'tags':src.tags(b)})
        rr.append(d)
save('rasters.json',rr)
with zipfile.ZipFile(ROOT/'Capas.qgz') as z:
    qgs=ET.fromstring(z.read(next(n for n in z.namelist() if n.endswith('.qgs'))))
    layers=[{'name':l.findtext('layername'),'source':l.findtext('datasource'),'provider':l.findtext('provider'),'crs':l.findtext('srs/spatialrefsys/authid')} for l in qgs.findall('.//projectlayers/maplayer')]
    save('qgis_layers.json',layers)
print('Auditoría guardada en',OUT)
