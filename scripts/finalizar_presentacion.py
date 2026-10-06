from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
pptx_path = ROOT / "GEOAI_Oro_Presentacion_Profesional.pptx"

def finalize():
    prs = Presentation(str(pptx_path))
    
    # 1. Slide 11: Añadir imagen del mapa comparativo
    slide_11 = prs.slides[10]
    img_path = ROOT / "reports/experimento_600_indicios/comparativa_mapa_v1_vs_experimental.png"
    if img_path.exists():
        # Comprobar si ya existe alguna imagen para no duplicar
        has_pic = any(s.shape_type == 13 for s in slide_11.shapes)
        if not has_pic:
            pic = slide_11.shapes.add_picture(str(img_path), 603504, 2240280, width=7315200)
            print(f"[OK] Mapa comparativo insertado en Diapositiva 11 (Alto: {pic.height})")
    
    # 2. Slide 16: Asegurar que Shape 10 tenga el nombre completo
    slide_16 = prs.slides[15]
    for s_idx, s in enumerate(slide_16.shapes):
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                if p.text.strip() == "GEOAI v1":
                    p.text = "GEOAI v1.0 & Exp. 787"
                    print(f"[OK] Diapositiva 16 Shape {s_idx} actualizado a 'GEOAI v1.0 & Exp. 787'")

    prs.save(str(pptx_path))
    print("[ÉXITO] Presentación finalizada con éxito.")

if __name__ == "__main__":
    finalize()
