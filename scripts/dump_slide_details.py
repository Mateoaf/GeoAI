from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
prs = Presentation(str(ROOT / "GEOAI_Oro_Presentacion_Profesional.pptx"))
slides_to_check = [3, 4, 5, 7, 8, 9, 10, 11, 13, 15]

lines = []
for idx in slides_to_check:
    slide = prs.slides[idx]
    lines.append(f"**************** SLIDE {idx+1} ****************")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            txts = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
            if txts:
                lines.append(f"Shape {s_idx} ({shape.name}): {' /// '.join(txts)}")
        elif shape.has_table:
            lines.append(f"Shape {s_idx} ({shape.name}) [TABLE]:")
            for r in shape.table.rows:
                lines.append("   | " + " | ".join(c.text.strip() for c in r.cells) + " |")
        elif shape.shape_type == 13: # picture
            lines.append(f"Shape {s_idx} ({shape.name}) [PICTURE]")

out_file = ROOT / "reports/target_slides_dump.txt"
out_file.write_text("\n".join(lines), encoding="utf-8")
print(f"Dump saved to {out_file}")
