"""
scripts/inspect_presentation.py
Extrae todo el contenido de las diapositivas de la presentación PPTX
para diagnosticar qué datos ficticios/desactualizados contiene.
"""
import sys
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
pptx_path = ROOT / "GEOAI_Oro_Presentacion_Profesional.pptx"

def inspect():
    prs = Presentation(str(pptx_path))
    output_lines = [f"Total diapositivas: {len(prs.slides)}\n"]

    for idx, slide in enumerate(prs.slides):
        title = slide.shapes.title.text if slide.shapes.title else "(Sin Título)"
        output_lines.append(f"=== DIAPOSITIVA {idx+1}: {title} ===")
        
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t and t != title:
                        output_lines.append(f"  - {t}")
            elif shape.has_table:
                output_lines.append("  [TABLA]:")
                for row in shape.table.rows:
                    row_txt = [cell.text.replace("\n", " ").strip() for cell in row.cells]
                    output_lines.append(f"    | {' | '.join(row_txt)} |")
        output_lines.append("")

    out_file = ROOT / "reports/inspeccion_presentacion.txt"
    out_file.parent.mkdir(parents=True, exist_ok=True)
    out_file.write_text("\n".join(output_lines), encoding="utf-8")
    print(f"Inspección guardada en: {out_file.relative_to(ROOT)}")

if __name__ == "__main__":
    inspect()
